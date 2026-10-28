"""Compare one batter's numbers across splits: a host country, an opponent, and
individual bowlers. Writes an Excel workbook where every stat is a formula over
the ball-by-ball data on the "Balls" sheet.

    python cricket/batter_splits.py --player "V Kohli" --team India \
        --data cricket/data/odis \
        --host-label "New Zealand" --host-cities Auckland Hamilton "Mount Maunganui" \
            Napier Wellington Christchurch Dunedin Nelson Queenstown Whangarei Lincoln \
        --opponent "New Zealand" \
        --bowlers "JA Duffy" "MJ Henry" "KA Jamieson" "MJ Santner" "LH Ferguson" "KDC Clarke" \
        --out cricket/output/kohli_odi_nz.xlsx

Cricsheet doesn't record a match's country, so the host country is given as a list
of its cities. Names use Cricsheet's format (initials + surname).
"""

import argparse
import glob
from collections import defaultdict
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1C5CAB")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
BOWLER_WICKETS = {"bowled", "caught", "caught and bowled", "lbw", "stumped", "hit wicket"}
NOT_OUT = {"not out", "retired hurt", "retired not out"}
PHASES = {
    "ODI": [("Powerplay (1-10)", 1, 10), ("Middle (11-40)", 11, 40), ("Death (41-50)", 41, 50)],
    "T20": [("Powerplay (1-6)", 1, 6), ("Middle (7-15)", 7, 15), ("Death (16-20)", 16, 20)],
}


def load(data_dirs, player, team, host_cities):
    """Every ball the player faced and every innings they batted, all matches."""
    host = {c.lower() for c in host_cities or []}
    balls, innings, matches, fmt = [], [], [], None
    for data_dir in data_dirs:
        for path in sorted(glob.glob(str(Path(data_dir) / "*.json"))):
            match = json.loads(Path(path).read_text())
            info = match["info"]
            if info.get("gender") != "male" or player not in info["players"].get(team, []):
                continue
            fmt = fmt or info.get("match_type")
            mid, date = Path(path).stem, info["dates"][0]
            opponent = next(t for t in info["teams"] if t != team)
            city = info.get("city") or ""
            in_host = "Yes" if city.lower() in host else "No"
            base = {"match_id": mid, "date": date, "opponent": opponent,
                    "venue": info.get("venue", ""), "city": city, "in_host": in_host}
            matches.append(base)
            for inn in match["innings"]:
                if inn.get("super_over") or inn["team"] != team:
                    continue
                order, out_how, seen_bowlers = [], None, set()
                for over in inn["overs"]:
                    for ball in over["deliveries"]:
                        for name in (ball["batter"], ball["non_striker"]):
                            if name not in order:
                                order.append(name)
                        out_kind = next(
                            (w["kind"] for w in ball.get("wickets", []) if w["player_out"] == player),
                            None,
                        )
                        out_how = out_how or out_kind
                        if ball["batter"] != player:
                            continue
                        bowler = ball["bowler"]
                        balls.append({
                            **base,
                            "over": over["over"] + 1,
                            "bowler": bowler,
                            "runs": ball["runs"]["batter"],
                            "legal": "Yes" if "wides" not in ball.get("extras", {}) else "No",
                            "out_kind": out_kind or "",
                            "bowler_wicket": int(out_kind in BOWLER_WICKETS),
                            "first_vs_bowler": int(bowler not in seen_bowlers),
                        })
                        seen_bowlers.add(bowler)
                if player not in order:
                    continue  # did not bat
                innings.append({**base, "position": order.index(player) + 1,
                                "dismissal": out_how or "not out"})
    by_date = lambda rows: pd.DataFrame(rows).sort_values(["date", "match_id"], kind="stable")  # noqa: E731
    return by_date(balls), by_date(innings), by_date(matches), fmt or "ODI"


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _table(ws, headers, widths):
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def write_workbook(path, player, balls, innings, matches, fmt, splits, bowlers):
    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    phase_ws = wb.create_sheet("By Phase")
    bowl_ws = wb.create_sheet("vs Bowlers")
    inn_ws = wb.create_sheet("Innings")
    ball_ws = wb.create_sheet("Balls")

    # ---- Balls: the raw data ------------------------------------------------
    phases = PHASES.get(fmt, PHASES["ODI"])
    phase_of = lambda o: next((n for n, lo, hi in phases if lo <= o <= hi), "Other")  # noqa: E731
    _table(ball_ws, ["Match ID", "Date", "Opponent", "Venue", "City", "In host country",
                     "Over", "Phase", "Bowler", "Runs off bat", "Legal ball", "Out on this ball",
                     "Wicket to bowler", "First ball from this bowler"],
           [11, 11, 14, 28, 14, 9, 6, 16, 16, 8, 7, 12, 8, 10])
    for b in balls.itertuples():
        ball_ws.append([b.match_id, b.date, b.opponent, b.venue, b.city, b.in_host, b.over,
                        phase_of(b.over), b.bowler, b.runs, b.legal, b.out_kind,
                        b.bowler_wicket, b.first_vs_bowler])
    nb = len(balls) + 1
    B = lambda c: f"Balls!${c}$2:${c}${nb}"  # noqa: E731

    # ---- Innings: details as values, numbers summed from Balls ---------------
    _table(inn_ws, ["Match ID", "Date", "Opponent", "Venue", "City", "In host country",
                    "Batting position", "Runs", "Balls", "4s", "6s", "Strike rate",
                    "Dismissal", "Out (1 = yes)"],
           [11, 11, 14, 28, 14, 9, 9, 7, 7, 5, 5, 9, 16, 8])
    for r, i in enumerate(innings.itertuples(), start=2):
        m = f"{B('A')},$A{r}"
        inn_ws.append([
            i.match_id, i.date, i.opponent, i.venue, i.city, i.in_host, i.position,
            f"=SUMIFS({B('J')},{m})",
            f'=COUNTIFS({m},{B("K")},"Yes")',
            f"=COUNTIFS({m},{B('J')},4)",
            f"=COUNTIFS({m},{B('J')},6)",
            f'=IF(I{r}=0,"",H{r}/I{r}*100)',
            i.dismissal,
            f'=IF(OR(M{r}="not out",M{r}="retired hurt",M{r}="retired not out"),0,1)',
        ])
        inn_ws.cell(row=r, column=12).number_format = "0.0"
    ni = len(innings) + 1
    I = lambda c: f"Innings!${c}$2:${c}${ni}"  # noqa: E731

    # ---- Summary: one column per split --------------------------------------
    # Each split is (label, (Balls col, value), (Innings col, value)); "*" = any text.
    summ["A1"] = f"{player} — {fmt} batting splits"
    summ["A1"].font = Font(name=FONT, bold=True, size=13)
    summ["A2"] = "Source: Cricsheet ball-by-ball data (cricsheet.org), men's matches."
    summ["A2"].font = Font(name=FONT, italic=True, color="555555")
    head = 4
    summ.cell(row=head, column=1, value="Stat")
    for c, (label, *_rest) in enumerate(splits, start=2):
        summ.cell(row=head, column=c, value=label)
    summ.cell(row=head, column=len(splits) + 2, value="How it's calculated")
    _style_header(summ, head, len(splits) + 2)

    def summary_rows(bc, bv, ic, iv, col, match_count):
        bcrit = f'{B(bc)},"{bv}"'
        icrit = f'{I(ic)},"{iv}"'
        legal = f'{bcrit},{B("K")},"Yes"'
        x = lambda stat: f"{col}{row_of[stat]}"  # noqa: E731
        return [
            ("Matches", match_count, "Matches in the XI (incl. did not bat)", None),
            ("Innings", f"=COUNTIFS({icrit})", "Innings batted", None),
            ("Not outs", f"={x('Innings')}-{x('Dismissals')}", "Innings − dismissals", None),
            ("Runs", f"=SUMIFS({I('H')},{icrit})", "Runs off the bat", None),
            ("Balls faced", f"=COUNTIFS({legal})", "Legal balls (wides excluded)", None),
            ("Strike rate", f'=IF({x("Balls faced")}=0,"",{x("Runs")}/{x("Balls faced")}*100)',
             "Runs ÷ balls × 100", "0.0"),
            ("Average", f'=IF({x("Dismissals")}=0,"",{x("Runs")}/{x("Dismissals")})',
             "Runs ÷ dismissals", "0.00"),
            ("Highest score",
             f'=_xlfn.MAXIFS({I("H")},{icrit})&IF(COUNTIFS({icrit},{I("H")},'
             f'_xlfn.MAXIFS({I("H")},{icrit}),{I("N")},0)>0,"*","")',
             "* = not out", None),
            ("100s", f'=COUNTIFS({icrit},{I("H")},">=100")', "", None),
            ("50s", f'=COUNTIFS({icrit},{I("H")},">=50",{I("H")},"<100")', "50 to 99", None),
            ("Ducks", f'=COUNTIFS({icrit},{I("H")},0,{I("N")},1)', "Out for 0", None),
            ("Fours", f"=COUNTIFS({bcrit},{B('J')},4)", "", None),
            ("Sixes", f"=COUNTIFS({bcrit},{B('J')},6)", "", None),
            ("Boundary % (4s+6s)",
             f'=IF({x("Balls faced")}=0,"",({x("Fours")}+{x("Sixes")})/{x("Balls faced")})',
             "(4s + 6s) ÷ balls", "0.0%"),
            ("Runs from boundaries %",
             f'=IF({x("Runs")}=0,"",(4*{x("Fours")}+6*{x("Sixes")})/{x("Runs")})',
             "Boundary runs ÷ runs", "0.0%"),
            ("Balls per boundary",
             f'=IF({x("Fours")}+{x("Sixes")}=0,"",{x("Balls faced")}/({x("Fours")}+{x("Sixes")}))',
             "", "0.0"),
            ("Dot balls", f"=COUNTIFS({legal},{B('J')},0)", "Legal balls, no runs off the bat", None),
            ("Dot ball %", f'=IF({x("Balls faced")}=0,"",{x("Dot balls")}/{x("Balls faced")})',
             "Dot balls ÷ balls", "0.0%"),
            ("Singles", f"=COUNTIFS({legal},{B('J')},1)", "", None),
            ("Rotation % (1-3 runs)",
             f'=IF({x("Balls faced")}=0,"",COUNTIFS({legal},{B("J")},">=1",{B("J")},"<=3")'
             f'/{x("Balls faced")})',
             "Balls scored 1, 2 or 3 off ÷ balls", "0.0%"),
            ("Balls per innings", f'=IF({x("Innings")}=0,"",{x("Balls faced")}/{x("Innings")})',
             "", "0.0"),
            ("Dismissals", f"=SUMIFS({I('N')},{icrit})", "", None),
            ("Balls per dismissal",
             f'=IF({x("Dismissals")}=0,"",{x("Balls faced")}/{x("Dismissals")})', "", "0.0"),
        ]

    row_of = defaultdict(int)  # placeholder so the first pass can list the labels
    labels = [r[0] for r in summary_rows("A", "*", "A", "*", "B", 0)]
    row_of = {label: head + 1 + k for k, label in enumerate(labels)}
    for c, (label, (bc, bv), (ic, iv), match_count) in enumerate(splits, start=2):
        col = get_column_letter(c)
        for stat, value, note, fmt_ in summary_rows(bc, bv, ic, iv, col, match_count):
            r = row_of[stat]
            summ.cell(row=r, column=1, value=stat)
            cell = summ.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(horizontal="right")
            if fmt_:
                cell.number_format = fmt_
            if stat == "Matches":
                cell.font = Font(name=FONT, color="0000FF")
                cell.fill = INPUT_FILL
            summ.cell(row=r, column=len(splits) + 2, value=note)
    summ.cell(row=row_of["Matches"], column=2).comment = Comment(
        "Counted from Cricsheet line-ups: matches in the playing XI, including any "
        "where the player did not bat.", "Claude")
    summ.column_dimensions["A"].width = 24
    for c in range(2, len(splits) + 2):
        summ.column_dimensions[get_column_letter(c)].width = 16
    summ.column_dimensions[get_column_letter(len(splits) + 2)].width = 40
    summ.freeze_panes = summ.cell(row=head + 1, column=2)

    # ---- By Phase: runs/balls/SR/dots per phase for each split ----------------
    cols = ["Split", "Phase", "Runs", "Balls", "Strike rate", "4s", "6s", "Dot ball %",
            "Boundary %", "Dismissals (on strike)"]
    _table(phase_ws, cols, [24, 18, 8, 8, 10, 6, 6, 10, 10, 10])
    r = 2
    for label, (bc, bv), _, _ in splits:
        for name, _, _ in phases:
            crit = f'{B(bc)},"{bv}",{B("H")},$B{r}'
            phase_ws.append([
                label, name,
                f"=SUMIFS({B('J')},{crit})",
                f'=COUNTIFS({crit},{B("K")},"Yes")',
                f'=IF(D{r}=0,"",C{r}/D{r}*100)',
                f"=COUNTIFS({crit},{B('J')},4)",
                f"=COUNTIFS({crit},{B('J')},6)",
                f'=IF(D{r}=0,"",COUNTIFS({crit},{B("K")},"Yes",{B("J")},0)/D{r})',
                f'=IF(D{r}=0,"",(F{r}+G{r})/D{r})',
                f'=COUNTIFS({crit},{B("L")},"?*")',
            ])
            phase_ws.cell(row=r, column=5).number_format = "0.0"
            phase_ws.cell(row=r, column=8).number_format = "0.0%"
            phase_ws.cell(row=r, column=9).number_format = "0.0%"
            r += 1

    # ---- vs Bowlers: one column per bowler ----------------------------------
    bowl_ws["A1"] = f"{player} vs individual bowlers ({fmt}s, all venues)"
    bowl_ws["A1"].font = Font(name=FONT, bold=True, size=13)
    bowl_ws["A2"] = ("Dismissals count only wickets credited to the bowler "
                     "(not run outs). Innings = innings in which he faced the bowler.")
    bowl_ws["A2"].font = Font(name=FONT, italic=True, color="555555")
    head = 4
    bowl_ws.cell(row=head, column=1, value="Stat")
    for c, name in enumerate(bowlers, start=2):
        bowl_ws.cell(row=head, column=c, value=name)
    _style_header(bowl_ws, head, len(bowlers) + 1)
    bstats = ["Innings faced", "Runs", "Balls", "Strike rate", "Dismissals", "Average",
              "Balls per dismissal", "Fours", "Sixes", "Boundary % (4s+6s)",
              "Runs from boundaries %", "Dot balls", "Dot ball %", "Singles",
              "Rotation % (1-3 runs)", "Caught", "Bowled", "LBW", "Other dismissals"]
    brow = {s: head + 1 + k for k, s in enumerate(bstats)}
    for c, name in enumerate(bowlers, start=2):
        col = get_column_letter(c)
        bc = f'{B("I")},{col}${head}'
        legal = f'{bc},{B("K")},"Yes"'
        x = lambda s: f"{col}{brow[s]}"  # noqa: E731
        values = {
            "Innings faced": f"=SUMIFS({B('N')},{bc})",
            "Runs": f"=SUMIFS({B('J')},{bc})",
            "Balls": f"=COUNTIFS({legal})",
            "Strike rate": f'=IF({x("Balls")}=0,"",{x("Runs")}/{x("Balls")}*100)',
            "Dismissals": f"=SUMIFS({B('M')},{bc})",
            "Average": f'=IF({x("Dismissals")}=0,"",{x("Runs")}/{x("Dismissals")})',
            "Balls per dismissal": f'=IF({x("Dismissals")}=0,"",{x("Balls")}/{x("Dismissals")})',
            "Fours": f"=COUNTIFS({bc},{B('J')},4)",
            "Sixes": f"=COUNTIFS({bc},{B('J')},6)",
            "Boundary % (4s+6s)": f'=IF({x("Balls")}=0,"",({x("Fours")}+{x("Sixes")})/{x("Balls")})',
            "Runs from boundaries %":
                f'=IF({x("Runs")}=0,"",(4*{x("Fours")}+6*{x("Sixes")})/{x("Runs")})',
            "Dot balls": f"=COUNTIFS({legal},{B('J')},0)",
            "Dot ball %": f'=IF({x("Balls")}=0,"",{x("Dot balls")}/{x("Balls")})',
            "Singles": f"=COUNTIFS({legal},{B('J')},1)",
            "Rotation % (1-3 runs)":
                f'=IF({x("Balls")}=0,"",COUNTIFS({legal},{B("J")},">=1",{B("J")},"<=3")/{x("Balls")})',
            "Caught": f'=COUNTIFS({bc},{B("L")},"caught")+COUNTIFS({bc},{B("L")},"caught and bowled")',
            "Bowled": f'=COUNTIFS({bc},{B("L")},"bowled")',
            "LBW": f'=COUNTIFS({bc},{B("L")},"lbw")',
            "Other dismissals": f'={x("Dismissals")}-{x("Caught")}-{x("Bowled")}-{x("LBW")}',
        }
        fmts = {"Strike rate": "0.0", "Average": "0.00", "Balls per dismissal": "0.0",
                "Boundary % (4s+6s)": "0.0%", "Runs from boundaries %": "0.0%",
                "Dot ball %": "0.0%", "Rotation % (1-3 runs)": "0.0%"}
        for s in bstats:
            bowl_ws.cell(row=brow[s], column=1, value=s)
            cell = bowl_ws.cell(row=brow[s], column=c, value=values[s])
            cell.alignment = Alignment(horizontal="right")
            if s in fmts:
                cell.number_format = fmts[s]
    bowl_ws.column_dimensions["A"].width = 24
    for c in range(2, len(bowlers) + 2):
        bowl_ws.column_dimensions[get_column_letter(c)].width = 13

    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font.name != FONT:
                    f = cell.font
                    cell.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color,
                                     size=f.size)
    for ws in (inn_ws, ball_ws):
        ws.auto_filter.ref = ws.dimensions
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--player", required=True)
    parser.add_argument("--team", required=True)
    parser.add_argument("--data", nargs="+", required=True)
    parser.add_argument("--host-label", help="name for the host-country split")
    parser.add_argument("--host-cities", nargs="+", default=[])
    parser.add_argument("--opponent")
    parser.add_argument("--bowlers", nargs="+", default=[])
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    balls, innings, matches, fmt = load(args.data, args.player, args.team, args.host_cities)
    if innings.empty:
        raise SystemExit("No innings found - check the player and team names.")
    splits = []
    if args.host_cities:
        splits.append((f"In {args.host_label or 'host country'}", ("F", "Yes"), ("F", "Yes"),
                       int((matches.in_host == "Yes").sum())))
    if args.opponent:
        splits.append((f"vs {args.opponent} (all venues)", ("C", args.opponent),
                       ("C", args.opponent), int((matches.opponent == args.opponent).sum())))
    splits.append((f"All {fmt}s (career)", ("A", "*"), ("A", "*"), len(matches)))

    missing = [b for b in args.bowlers if b not in set(balls.bowler)]
    if missing:
        print("Never faced (check spelling):", ", ".join(missing))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(args.out, args.player, balls, innings, matches, fmt, splits, args.bowlers)
    print(f"{len(matches)} matches, {len(innings)} innings, {len(balls)} balls -> {args.out}")


if __name__ == "__main__":
    main()
