"""Export every practice lap (sector times and average speeds) to an Excel workbook.

Sheet "All Laps" has one row per lap; sheet "Averages" averages each driver's
representative laps per session. Speeds are Excel formulas (sector length / time),
so they update if you change the track lengths in the "All Laps" input cells.

Sector lengths are not published by F1, so they are estimated from car telemetry:
speed is integrated over the fastest laps of the weekend to get each sector's share
of the lap, which is then scaled to the official lap length passed with --lap-km.

Usage:
    python f1/practice_laps_to_excel.py --year 2026 --location Austin --lap-km 5.513
"""

import argparse
import hashlib
import json

import numpy as np
import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from practice_race_pace import API, CACHE, HERE, PUSH_LAP_FACTOR, SECTORS, fetch

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1C5CAB")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")

COLUMNS = [
    "Driver name",
    "Session",
    "Sector 1 time (s)",
    "Sector 2 time (s)",
    "Sector 3 time (s)",
    "Lap time (s)",
    "Sector 1 speed (km/h)",
    "Sector 2 speed (km/h)",
    "Sector 3 speed (km/h)",
    "Lap speed (km/h)",
    "Team",
]
LAP_EXTRA = ["Lap number", "Tyre", "Tyre age (laps)", "Lap type", "Used in average"]
AVG_EXTRA = ["Laps averaged"]


def fetch_car_data(session_key, driver_number, start, end):
    """Car telemetry (speed etc.) for one driver between two timestamps, cached."""
    query = (
        f"session_key={session_key}&driver_number={driver_number}"
        f"&date>={start.isoformat()}&date<={end.isoformat()}"
    )
    path = CACHE / f"car_data_{hashlib.md5(query.encode()).hexdigest()}.json"
    if path.exists():
        return json.loads(path.read_text())
    resp = requests.get(f"{API}/car_data?{query}", timeout=60)
    resp.raise_for_status()
    CACHE.mkdir(parents=True, exist_ok=True)
    path.write_text(resp.text)
    return resp.json()


def sector_shares(laps, n_laps=10):
    """Estimate each sector's share of the lap distance from speed telemetry."""
    timed = laps.dropna(subset=["lap_duration", "date_start"] + SECTORS)
    timed = timed[~timed.is_pit_out_lap.fillna(False).astype(bool)]
    fastest = timed.sort_values("lap_duration").drop_duplicates(["session_key", "driver_number"])

    shares = []
    for lap in fastest.head(n_laps).itertuples():
        t0 = pd.Timestamp(lap.date_start)
        pad = pd.Timedelta(seconds=2)
        car = pd.DataFrame(
            fetch_car_data(lap.session_key, lap.driver_number, t0 - pad,
                           t0 + pd.Timedelta(seconds=lap.lap_duration) + pad)
        )
        if car.empty:
            continue
        t = pd.to_datetime(car.date, format="ISO8601").astype("int64").to_numpy() / 1e9
        order = np.argsort(t)
        t, v = t[order], car.speed.to_numpy()[order] / 3.6
        dist = np.concatenate([[0], np.cumsum((v[1:] + v[:-1]) / 2 * np.diff(t))])
        start = t0.value / 1e9
        bounds = start + np.cumsum(
            [0, lap.duration_sector_1, lap.duration_sector_2, lap.duration_sector_3]
        )
        d = np.diff(np.interp(bounds, t, dist))
        shares.append(d / d.sum())
    if not shares:
        raise SystemExit("No telemetry available to estimate sector lengths.")
    return np.mean(shares, axis=0), len(shares)


def load_laps(year, location, sessions):
    """All laps for the chosen sessions, with driver, tyre and lap-type info."""
    frames = []
    for s in fetch("sessions", year=year, location=location):
        if s["session_name"] not in sessions:
            continue
        key = s["session_key"]
        laps = pd.DataFrame(fetch("laps", session_key=key))
        if laps.empty:
            print(f"{s['session_name']}: no lap data yet, skipped")
            continue
        stints = pd.DataFrame(fetch("stints", session_key=key))
        drivers = pd.DataFrame(fetch("drivers", session_key=key))

        laps = laps.merge(stints, on=["driver_number", "session_key", "meeting_key"], how="left")
        in_stint = laps.lap_number.between(laps.lap_start, laps.lap_end)
        # Keep the matching stint for each lap, or one stint-less row if none matches.
        laps = pd.concat([
            laps[in_stint],
            laps[~in_stint].drop_duplicates(["driver_number", "lap_number"])
            .assign(compound=None, tyre_age_at_start=np.nan, lap_start=np.nan, lap_end=np.nan),
        ]).drop_duplicates(["driver_number", "lap_number"])

        laps["tyre_age"] = laps.tyre_age_at_start + laps.lap_number - laps.lap_start
        laps["lap_type"] = np.select(
            [laps.is_pit_out_lap.fillna(False).astype(bool), laps.lap_number == laps.lap_end],
            ["Out lap", "In lap"],
            "Timed",
        )
        drivers["driver_name"] = (
            drivers.first_name.fillna("") + " " + drivers.last_name.fillna("")
        ).str.strip().where(drivers.last_name.notna(), drivers.full_name)
        laps = laps.merge(drivers[["driver_number", "driver_name", "team_name"]], on="driver_number")
        laps["session"] = s["session_name"]
        frames.append(laps)
    if not frames:
        raise SystemExit("No lap data found for those sessions.")
    return pd.concat(frames).sort_values(["session", "driver_name", "lap_number"])


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"


def build_workbook(laps, sessions, lap_km, shares, n_share_laps, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "All Laps"
    headers = COLUMNS + LAP_EXTRA
    ws.append(headers)
    n = len(laps)
    last = n + 1

    # Inputs block to the right of the lap table.
    ic, vc = len(headers) + 2, len(headers) + 3  # label column, value column
    L, V = get_column_letter(ic), get_column_letter(vc)
    inputs = [
        ("Track inputs", None),
        ("Lap length (km)", lap_km),
        ("Sector 1 length (km)", round(lap_km * shares[0], 3)),
        ("Sector 2 length (km)", round(lap_km * shares[1], 3)),
        ("Sector 3 length (km)", None),
        ("Average uses laps within this × session fastest", PUSH_LAP_FACTOR),
        ("Session fastest lap (s)", None),
    ]
    for i, (label, value) in enumerate(inputs, start=1):
        ws.cell(row=i, column=ic, value=label).font = Font(name=FONT, bold=value is None)
        if value is not None:
            cell = ws.cell(row=i, column=vc, value=value)
            cell.font = Font(name=FONT, color="0000FF")
            cell.fill = INPUT_FILL
    # Sector 3 is the remainder, so the three sectors always add up to the lap.
    ws[f"{V}5"] = f"={V}2-{V}3-{V}4"
    ws[f"{V}5"].font = Font(name=FONT)
    ws[f"{V}6"].number_format = "0.00"
    for r in (2, 3, 4, 5):
        ws[f"{V}{r}"].number_format = "0.000"
    ws[f"{V}2"].comment = Comment("Official circuit length, from the F1 circuit guide.", "Claude")
    ws[f"{V}3"].comment = Comment(
        f"Estimated: F1 does not publish sector lengths. Each sector's share of the lap "
        f"was measured by integrating speed telemetry (OpenF1 car_data) over the "
        f"{n_share_laps} fastest laps of the weekend, then scaled to the lap length.",
        "Claude",
    )
    fastest_rows = {}
    for i, sess in enumerate(sessions, start=8):
        ws.cell(row=i, column=ic, value=sess).font = Font(name=FONT)
        c = ws.cell(row=i, column=vc, value=f'=_xlfn.MINIFS($F$2:$F${last},$B$2:$B${last},{L}{i})')
        c.font, c.number_format = Font(name=FONT), "0.000"
        fastest_rows[sess] = i
    f_first, f_last = 8, 7 + len(sessions)
    notes_row = f_last + 2
    notes = [
        "Notes",
        "Times from OpenF1 (openf1.org). Speeds = sector or lap length ÷ time × 3600.",
        "Lap type: Out lap = leaving the pits, In lap = last lap of a tyre stint (into the pits).",
        "Used in average = Timed lap with all sectors, within the factor above of its "
        "session's fastest lap (drops cool-down and slow laps).",
        "Blue cells on yellow are inputs you can change; everything else recalculates.",
    ]
    for i, text in enumerate(notes):
        ws.cell(row=notes_row + i, column=ic, value=text).font = Font(name=FONT, bold=i == 0)

    for r, lap in enumerate(laps.itertuples(), start=2):
        vals = [
            lap.driver_name, lap.session,
            lap.duration_sector_1, lap.duration_sector_2, lap.duration_sector_3, lap.lap_duration,
        ]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=None if pd.isna(v) else v)
        for c, (t_col, len_cell) in enumerate(
            [("C", f"${V}$3"), ("D", f"${V}$4"), ("E", f"${V}$5"), ("F", f"${V}$2")], start=7
        ):
            ws.cell(row=r, column=c, value=f'=IF({t_col}{r}="","",{len_cell}/{t_col}{r}*3600)')
        ws.cell(row=r, column=11, value=lap.team_name)
        ws.cell(row=r, column=12, value=int(lap.lap_number))
        ws.cell(row=r, column=13, value=None if pd.isna(lap.compound) else str(lap.compound).title())
        ws.cell(row=r, column=14, value=None if pd.isna(lap.tyre_age) else int(lap.tyre_age))
        ws.cell(row=r, column=15, value=lap.lap_type)
        ws.cell(
            row=r, column=16,
            value=(
                f'=IF(AND(O{r}="Timed",COUNT(C{r}:F{r})=4,'
                f'F{r}<=INDEX(${V}${f_first}:${V}${f_last},MATCH(B{r},${L}${f_first}:${L}${f_last},0))'
                f'*${V}$6),"Yes","No")'
            ),
        )

    # Averages sheet: one row per driver per session, sorted by average lap time.
    wa = wb.create_sheet("Averages")
    wa.append(COLUMNS + AVG_EXTRA)
    laps = laps.assign(
        _ok=(laps.lap_type == "Timed")
        & laps[SECTORS + ["lap_duration"]].notna().all(axis=1)
        & (laps.lap_duration <= laps.groupby("session").lap_duration.transform("min") * PUSH_LAP_FACTOR)
    )
    order = (
        laps[laps._ok].groupby(["session", "driver_name"]).lap_duration.mean()
        .rename("avg").reset_index()
    )
    drivers = laps.groupby(["session", "driver_name"]).team_name.first().reset_index()
    drivers = drivers.merge(order, how="left")
    drivers["sess_order"] = drivers.session.map({s: i for i, s in enumerate(sessions)})
    drivers = drivers.sort_values(["sess_order", "avg"], na_position="last")

    rng = lambda col: f"'All Laps'!${col}$2:${col}${last}"  # noqa: E731
    crit = lambda r: f'{rng("A")},$A{r},{rng("B")},$B{r},{rng("P")},"Yes"'  # noqa: E731
    for r, d in enumerate(drivers.itertuples(), start=2):
        wa.cell(row=r, column=1, value=d.driver_name)
        wa.cell(row=r, column=2, value=d.session)
        for c, col in enumerate("CDEF", start=3):
            wa.cell(row=r, column=c, value=f'=IFERROR(AVERAGEIFS({rng(col)},{crit(r)}),"")')
        for c, (t_col, len_cell) in enumerate(
            [("C", "$3"), ("D", "$4"), ("E", "$5"), ("F", "$2")], start=7
        ):
            ref = f"'All Laps'!${V}{len_cell}"
            wa.cell(row=r, column=c, value=f'=IF({t_col}{r}="","",{ref}/{t_col}{r}*3600)')
        wa.cell(row=r, column=11, value=d.team_name)
        wa.cell(row=r, column=12, value=f"=COUNTIFS({crit(r)})")

    widths = {1: 22, 2: 12, 11: 18, 15: 10, 16: 11}
    for sheet, rows in ((ws, last), (wa, len(drivers) + 1)):
        style_header(sheet, sheet.max_column if sheet is wa else len(headers))
        for row in sheet.iter_rows(min_row=2, max_row=rows, max_col=len(COLUMNS) + 5):
            for cell in row:
                cell.font = Font(name=FONT)
                if 3 <= cell.column <= 6:
                    cell.number_format = "0.000"
                elif 7 <= cell.column <= 10:
                    cell.number_format = "0.0"
        for c in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(c)].width = widths.get(c, 13)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column if sheet is wa else len(headers))}{rows}"
    ws.column_dimensions[L].width = 44
    ws.column_dimensions[V].width = 12
    thin = Side(style="thin", color="BFBFBF")
    for r in range(2, 7):
        ws[f"{V}{r}"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb.calculation.fullCalcOnLoad = True  # openpyxl stores no results; Excel computes them
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--location", required=True, help='e.g. "Austin", "Monza"')
    parser.add_argument("--lap-km", type=float, required=True, help="official lap length")
    parser.add_argument("--sessions", nargs="+", default=["Practice 1", "Practice 2"])
    parser.add_argument("--out", default=None, help="output .xlsx path")
    args = parser.parse_args()

    laps = load_laps(args.year, args.location, args.sessions)
    sessions = [s for s in args.sessions if s in set(laps.session)]
    shares, n_share_laps = sector_shares(laps)
    out = args.out or HERE / "output" / f"{args.year}_{args.location.lower()}_practice_laps.xlsx"
    (HERE / "output").mkdir(exist_ok=True)
    build_workbook(laps, sessions, args.lap_km, shares, n_share_laps, out)
    print(f"{len(laps)} laps, sessions {sessions}")
    print("Sector shares:", np.round(shares, 4), f"from {n_share_laps} laps")
    print(f"Saved {out}")


if __name__ == "__main__":
    main()
